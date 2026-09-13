# -*- coding: utf-8 -*-
"""Test de mutation : que surveille reellement la batterie de controles.

Un modele declare ses propres controles. Il ne se surveille donc que la ou il
a pense a se surveiller, et personne ne sait quelle part du classeur echappe a
cette surveillance. Ce module la mesure au lieu de la supposer.

Le principe est brutal et sans appel : on **casse** une cellule - on remplace
sa formule par sa valeur multipliee par 1,01 - on recalcule, et on regarde si
un seul controle du modele bronche. Si rien ne bouge, la corruption est
INVISIBLE : cette cellule est hors de portee de la batterie.

Repete sur un echantillon stratifie, ca donne un chiffre qu'aucune autre
methode ne donne : le taux de cellules dont on peut fausser la valeur sans que
le modele s'en apercoive. C'est la couverture reelle de l'organe de controle.

Une precision qui evite un contresens : une mutation non detectee n'est pas
forcement un defaut. Une cellule de presentation n'a pas a etre surveillee. Ce
qui compte est la ou tombent les angles morts - si ce sont les feuilles qui
portent le livrable, la batterie decore.

Chaque mutation coute un recalcul Excel complet. Le module travaille sur une
copie, une mutation a la fois - autrement on ne saurait pas a qui attribuer le
mouvement - et il annonce toujours la taille de son echantillon face a la
population.

Usage :
    python mutation.py "Classeur.xlsx" --batterie "22. Checks!D51" --echantillon 40
    python mutation.py "Classeur.xlsx" --batterie "..." --feuilles "08.,09." --json m.json
"""
import argparse
import collections
import json
import os
import random
import re
import sys
import time

import noyau
from openpyxl.utils import get_column_letter
import rejouer


def sentinelles(C, batterie):
    """Les cellules que l'on surveillera apres chaque mutation : le total de
    la batterie, plus chaque ligne de controle qu'on sait reperer."""
    out = [batterie] if batterie else []
    for (f, l, c), v in C.textes.items():
        if not isinstance(v, str):
            continue
        s = v.lower()
        if any(m in s for m in ("violation", "must be nil", "controlled to nil",
                                "residual", "tie -")):
            # la ligne porte un controle : on prend sa premiere colonne calculee
            for cc in range(c + 1, c + 14):
                if (f, l, cc) in C.formules:
                    out.append(f"'{f}'!{get_column_letter(cc)}{l}")
                    break
    # dedoublonne en gardant l'ordre
    vus = set()
    return [x for x in out if not (x in vus or vus.add(x))]


def couverture_structurelle(C, sentinelles):
    """Quelles cellules PEUVENT, en principe, deplacer un controle.

    Le graphe repond gratuitement a une moitie de la question, et sur 100 % du
    classeur : si aucune sentinelle n'est descendante d'une cellule, alors la
    fausser ne peut RIEN deplacer. C'est une preuve, pas un echantillon.

    On remonte donc les precedents depuis chaque sentinelle : l'ensemble des
    ancetres est exactement l'ensemble des cellules surveillables. Tout ce qui
    est en dehors est un angle mort demontre, sans un seul recalcul.

    Ce que le graphe ne dit PAS : parmi les cellules surveillables, lesquelles
    voient leur effet s'annuler en route - le cas du tie tautologique, dont les
    deux membres bougent ensemble. Cela demande de recalculer, et c'est l'objet
    de l'echantillon.
    """
    adj = C.graphe()                       # adj[i] = precedents de i
    depart = []
    for s in sentinelles:
        try:
            feuille, ligne, col = _adresse(C, s)
        except Exception:
            continue
        i = C.idx.get((feuille, ligne, col))
        if i is not None:
            depart.append(i)
    vus = set(depart)
    pile = list(depart)
    while pile:
        i = pile.pop()
        for j in adj[i]:
            if j not in vus:
                vus.add(j)
                pile.append(j)
    return {C.rev_idx[i] for i in vus}


_ADR = re.compile(r"^'?([^'!]+)'?!\$?([A-Z]{1,3})\$?([0-9]+)$")   # dollars admis : l'autotest les exige


def _adresse(C, s):
    m = _ADR.match(s.strip())
    if not m:
        raise ValueError(s)
    from openpyxl.utils import column_index_from_string
    return m.group(1), int(m.group(3)), column_index_from_string(m.group(2))


def echantillon_stratifie(C, taille, feuilles=None, graine=12345, restreint=None):
    """Un tirage par feuille, proportionnel, pour ne pas mesurer une seule
    feuille en croyant mesurer le classeur."""
    par_feuille = collections.defaultdict(list)
    for cle in C.formules:
        if feuilles and not any(cle[0].startswith(p) for p in feuilles):
            continue
        if restreint is not None and cle not in restreint:
            continue
        # on ne mute que ce qui porte une valeur numerique : muter une cellule
        # vide ou textuelle ne teste rien
        v = C.valeurs.get(cle)
        if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v) > 1e-09:
            par_feuille[cle[0]].append(cle)
    total = sum(len(v) for v in par_feuille.values())
    if not total:
        return [], 0
    rnd = random.Random(graine)
    tirage = []
    for f, cells in sorted(par_feuille.items()):
        n = max(1, round(taille * len(cells) / total))
        tirage.extend(rnd.sample(cells, min(n, len(cells))))
    rnd.shuffle(tirage)
    return tirage[:taille], total


def autotest():
    """La mutation mesure ce que la BATTERIE surveille : si son adressage se
    casse, elle rend une couverture flatteuse sans avoir rien casse.

    On verifie la lecture d'adresse, dans ses trois ecritures - avec quotes,
    sans quotes, avec dollars - et le refus d'une adresse qui n'en est pas une.
    """
    bons = ["'01. Input_Sheet'!D12", "Modele!AA3", "'Deux mots'!$B$7"]
    for a in bons:
        if not _ADR.match(a):
            raise SystemExit(f"ARRET - la mutation ne sait plus lire l'adresse {a!r}")
    for a in ("D12", "pas une adresse", ""):
        if _ADR.match(a):
            raise SystemExit(f"ARRET - la mutation accepte {a!r} pour une adresse")


def main():
    autotest()
    ap = argparse.ArgumentParser()
    ap.add_argument("classeur")
    ap.add_argument("--batterie", help="adresse du total de la batterie ; trouvee toute seule si omise")
    ap.add_argument("--echantillon", type=int, default=30)
    ap.add_argument("--feuilles", help="prefixes de feuilles, separes par des virgules")
    ap.add_argument("--facteur", type=float, default=1.01,
                    help="facteur de corruption applique a la valeur (defaut +1 %%)")
    ap.add_argument("--iterations", type=int, default=500)
    ap.add_argument("--json")
    a = ap.parse_args()

    C = noyau.Classeur(a.classeur)
    if not a.batterie:
        a.batterie = noyau.trouver_batterie(C)
        if a.batterie:
            print(f"batterie trouvee automatiquement : {a.batterie}")
    feuilles = [x.strip() for x in a.feuilles.split(",")] if a.feuilles else None
    sent = sentinelles(C, a.batterie)
    surveillables = couverture_structurelle(C, sent)
    cibles, population = echantillon_stratifie(C, a.echantillon, feuilles,
                                               restreint=surveillables)

    total_formules = len(C.formules)
    n_surv = len(surveillables & set(C.formules))
    pct_struct = n_surv / max(total_formules, 1) * 100
    print(f"TEST DE MUTATION - {os.path.basename(a.classeur)}")
    print(f"sentinelles surveillees    : {len(sent)}")
    print()
    print("  1. COUVERTURE STRUCTURELLE - exhaustive, deduite du graphe")
    print(f"     cellules de formule                : {total_formules:>8,}")
    print(f"     dont une sentinelle est descendante: {n_surv:>8,}   {pct_struct:5.1f} %")
    print(f"     ANGLE MORT DEMONTRE                : "
          f"{total_formules - n_surv:>8,}   {100 - pct_struct:5.1f} %")
    print("     Ces cellules ne peuvent RIEN deplacer : aucun controle n'est en aval.")
    print("     C'est une preuve, pas un echantillon - aucun recalcul necessaire.")
    print()
    print("  2. TAUX DE PROPAGATION - echantillonne, parmi les surveillables")
    print(f"     population surveillable et mutable : {population:,}")
    print(f"     echantillon                        : {len(cibles)} "
          f"({len(cibles)/max(population,1)*100:.2f} %)")
    print(f"corruption appliquee       : valeur x {a.facteur}")
    print(f"cout                       : {len(cibles)} recalculs Excel\n")

    # etat de reference
    R0 = rejouer.Rejeu(a.classeur, etiquette="temoin")
    base = R0.relever(sent)
    R0.nettoyer()

    lignes = []
    invisibles = collections.Counter()
    vues = collections.Counter()
    for i, cle in enumerate(cibles, 1):
        t0 = time.time()
        adresse = f"'{cle[0]}'!{get_column_letter(cle[2])}{cle[1]}"
        valeur = C.valeurs[cle]
        R = rejouer.Rejeu(a.classeur, etiquette=f"mut{i}")
        try:
            R.poser({adresse: valeur * a.facteur})
            rec = R.recalculer(iterations=a.iterations)
            if rec["code"] != 0:
                etat, detectee = "recalcul en echec", None
            else:
                apres = R.relever(sent)
                bouge = [s for s in sent
                         if rejouer.ecart_relatif(base.get(s), apres.get(s)) > 1e-06
                         and abs((apres.get(s) or 0) - (base.get(s) or 0)) > 1e-06]
                detectee = bool(bouge)
                etat = f"{len(bouge)} sentinelle(s)" if bouge else "AUCUNE REACTION"
            if detectee is False:
                invisibles[cle[0]] += 1
            elif detectee:
                vues[cle[0]] += 1
            print(f"  [{i:>3}/{len(cibles)}] {'vue ' if detectee else 'INVISIBLE'} "
                  f"{adresse:<40} {etat:<22} {time.time()-t0:.0f}s")
            lignes.append({"cellule": adresse, "feuille": cle[0],
                           "valeur": valeur, "detectee": detectee, "etat": etat})
        finally:
            R.nettoyer()

    testees = sum(invisibles.values()) + sum(vues.values())
    taux = sum(vues.values()) / max(testees, 1) * 100
    print()
    print("=" * 74)
    reelle = pct_struct * taux / 100
    print(f"TAUX DE PROPAGATION (parmi les surveillables) : {taux:.1f} %")
    print(f"  {sum(vues.values())} corruptions detectees, "
          f"{sum(invisibles.values())} passees inapercues, sur {testees} testees")
    print()
    print(f"COUVERTURE REELLE = structurelle x propagation = "
          f"{pct_struct:.1f} % x {taux:.1f} % = {reelle:.1f} %")
    print("  Une corruption non detectee alors qu'une sentinelle est en aval")
    print("  signale un tie TAUTOLOGIQUE : ses deux membres bougent ensemble.")
    if invisibles:
        print("\n  ou sont les angles morts :")
        for f, n in invisibles.most_common():
            tot = n + vues.get(f, 0)
            print(f"      {f:<34} {n:>4} / {tot:<4} invisibles")
    print()
    print("  Rappel : une mutation non detectee n'est pas un defaut en soi - une")
    print("  cellule de presentation n'a pas a etre surveillee. Ce qui compte est")
    print("  la ou tombent les angles morts.")
    print(f"  Et l'echantillon ne couvre que {len(cibles)/max(population,1)*100:.2f} % "
          f"de la population : ce taux est une estimation, pas un recensement.")

    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump({"classeur": os.path.basename(a.classeur),
                       "population": population, "echantillon": len(cibles),
                       "sentinelles": len(sent),
                       "couverture_structurelle_pct": round(pct_struct, 1),
                       "cellules_surveillables": n_surv,
                       "angle_mort_demontre": total_formules - n_surv,
                       "taux_de_propagation_pct": round(taux, 1),
                       "couverture_reelle_pct": round(pct_struct * taux / 100, 1),
                       "angles_morts": dict(invisibles.most_common()),
                       "lignes": lignes}, fh, ensure_ascii=False, indent=2)
        print(f"\necrit : {a.json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

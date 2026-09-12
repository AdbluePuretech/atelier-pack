"""Compare le recalcul LibreOffice au cache Excel, cellule par cellule.

Un ecart ici ne se voit nulle part ailleurs : le classeur s'ouvre juste dans
Excel, la rubric le note a 100 %, et le service QA — qui recalcule sous
LibreOffice — le note faux. C'est exactement ce qui a coule le pack Final.
"""
import datetime
import json
import sys
from openpyxl import load_workbook

# LibreOffice rend une date en NUMERO DE SERIE, openpyxl en `datetime`. Sans
# cette conversion, les cent quatre-vingts cellules de la ligne de dates sont
# comptees « absentes » et ne sont jamais comparees — or une ligne de dates qui
# derive sous LibreOffice fait deriver tous les drapeaux de periode avec elle.
EPOQUE = datetime.datetime(1899, 12, 30)


def serie(x):
    if isinstance(x, datetime.datetime):
        return (x - EPOQUE).total_seconds() / 86400
    if isinstance(x, datetime.date):
        return (datetime.datetime(x.year, x.month, x.day) - EPOQUE).days
    return x

lo = json.load(open(sys.argv[2]))
wb = load_workbook(sys.argv[1], data_only=True)

compare = ecarts = absents = err = 0
pires, par_feuille = [], {}
for feuille, cells in lo["valeurs"].items():
    if feuille not in wb.sheetnames:
        continue
    ws = wb[feuille]
    for cle, v in cells.items():
        r, c = (int(x) for x in cle.split(":"))
        x = serie(ws.cell(row=r, column=c).value)
        if isinstance(v, str):
            err += 1
            pires.append((1e9, feuille, r, c, x, v))
            continue
        if not isinstance(x, (int, float)):
            absents += 1
            continue
        compare += 1
        d = abs(v - x) / max(abs(x), 1e-9)
        if d > 1e-6 and abs(v - x) > 1e-6:
            ecarts += 1
            par_feuille[feuille] = par_feuille.get(feuille, 0) + 1
            pires.append((d, feuille, r, c, x, v))

print(f"  cellules comparees        : {compare:,}")
print(f"  cellules en erreur (LO)   : {err}")
print(f"  numeriques cote LO mais pas cote Excel : {absents}")
print(f"  ECARTS > 1e-6 relatif     : {ecarts}")
if par_feuille:
    print("  par feuille :")
    for f, n in sorted(par_feuille.items(), key=lambda x: -x[1]):
        print(f"    {f:<26} {n}")
    print("  les dix plus gros :")
    for d, f, r, c, x, v in sorted(pires, reverse=True)[:10]:
        print(f"    {f}!r{r}c{c}  Excel {x!r:>20}  LibreOffice {v!r:>20}  ecart {d:.2%}")

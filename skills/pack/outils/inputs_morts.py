# -*- coding: utf-8 -*-
"""Les hypotheses de l'Input_Sheet que le modele ne consomme jamais.

Un input mort n'est pas du poids inutile, c'est un PIEGE INVOLONTAIRE. Le
candidat n'a que le prompt et l'input sheet ; il suppose, a juste titre, que
tout ce qu'on lui donne sert. Une ligne « Direct labour - Fire & Life Safety :
48% » laissee dans la feuille apres que le moteur de cout est passe a un
effectif multiplie par un salaire ne fait pas que traîner : elle DESIGNE la
methode qu'on veut precisement lui voir eviter. Le pack devient injuste, et
l'audit d'equite le renverra.

Le controle ne regarde pas les noms definis mais les FORMULES : une plage
nommee peut exister sans que rien ne la cite. On compte donc, pour chaque
cellule d'hypothese, les references qui la visent — par son nom defini comme
par son adresse.

    python inputs_morts.py <golden.xlsx> [feuille] [noms,toleres]

Les noms toleres sont les hypotheses qu'aucune formule ne peut citer par
construction — le reglage d'iteration du classeur, par exemple, que le
moteur lit dans calcPr et non dans une cellule.

Sortie : une ligne par hypothese morte, et un code de sortie non nul s'il y en a.
"""
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

CLASSEUR = sys.argv[1]
FEUILLE = sys.argv[2] if len(sys.argv) > 2 else "Input_Sheet"

wb = load_workbook(CLASSEUR, data_only=False)
if FEUILLE not in wb.sheetnames:
    raise SystemExit("  ARRET — feuille absente : " + FEUILLE)
ws = wb[FEUILLE]

# ---- 1. toutes les formules du classeur, hors la feuille d'hypotheses -------
# On exclut la feuille elle-meme : une hypothese citee UNIQUEMENT par une autre
# ligne de la meme feuille, elle-meme morte, reste morte.
formules = []
for f in wb.worksheets:
    for row in f.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formules.append((f.title, c.value))
hors = "\n".join(v for t, v in formules if t != FEUILLE)
dedans = "\n".join(v for t, v in formules if t == FEUILLE)

# ---- 2. les noms definis qui pointent sur une cellule de la feuille ---------
noms = {}
for n, d in wb.defined_names.items():
    m = re.match(r"^'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)$", str(d.value))
    if m and m.group(1) == FEUILLE:
        noms.setdefault((m.group(2), int(m.group(3))), []).append(n)


def cite(colonne, ligne, texte, local=False):
    """Une cellule est-elle visee dans ce texte, par nom ou par adresse ?

    `local` vaut pour les formules de la feuille elle-meme : elles citent leurs
    voisines SANS nommer la feuille (« =D$32 »). En exigeant le prefixe partout,
    on declare morte l'hypothese qu'une ligne de serie resout deux cents lignes
    plus bas — et c'est toute la chaine du CPI qui tombe.
    """
    for n in noms.get((colonne, ligne), []):
        if re.search(r"(?<![A-Za-z0-9_])" + re.escape(n) + r"(?![A-Za-z0-9_])", texte):
            return True
    prefixe = r"(?:{f}!)?".format(f=re.escape(FEUILLE)) if local \
        else r"{f}!".format(f=re.escape(FEUILLE))
    if re.search(r"(?<![A-Za-z0-9_!]){p}\$?{c}\$?{l}(?![0-9])".format(
            p=prefixe, c=colonne, l=ligne), texte):
        return True
    # une plage verticale qui englobe la ligne : Input_Sheet!$D$30:$D$40
    for m in re.finditer(prefixe + r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", texte):
        c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
        if c1 <= colonne <= c2 and r1 <= ligne <= r2:
            return True
    return False


TOLERES = set(sys.argv[3].split(",")) if len(sys.argv) > 3 else set()

# ---- 3. la vie se propage par relais ---------------------------------------
# Une hypothese peut n'etre citee par aucune feuille aval et rester vivante : la
# feuille porte souvent une ligne de SERIE qui la resout — « Live CPI, resolved
# through the case selector » — et c'est cette serie que le modele consomme. En
# s'arretant aux citations externes, on declarerait morte l'hypothese qui nourrit
# tout le modele. On part donc des lignes citees de l'exterieur et on remonte,
# jusqu'a ce que plus rien ne bouge.
formules_ligne = {}
for row in ws.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            formules_ligne.setdefault(c.row, []).append(c.value)

vivantes = set()
for r in range(1, ws.max_row + 1):
    if any(cite(get_column_letter(c), r, hors) for c in range(4, ws.max_column + 1)):
        vivantes.add(r)
while True:
    gagnees = set()
    for rv in vivantes:
        texte = "\n".join(formules_ligne.get(rv, []))
        if not texte:
            continue
        for r in range(1, ws.max_row + 1):
            if r in vivantes or r in gagnees:
                continue
            if any(cite(get_column_letter(c), r, texte, local=True)
                   for c in range(4, ws.max_column + 1)):
                gagnees.add(r)
    if not gagnees:
        break
    vivantes |= gagnees

morts, vivants, tolerees = [], 0, 0
for r in range(1, ws.max_row + 1):
    lab = ws.cell(row=r, column=3).value
    if not isinstance(lab, str) or not lab.strip():
        continue
    if re.match(r"^\s*\d+ - ", lab):          # bandeau de section
        continue
    # On balaye toute la LARGEUR de la ligne, pas les six colonnes de periode.
    # Le registre de contrats s'etale bien au-dela : la colonne du taux de
    # conversion du cas 2, celle du master service agreement et celle de la
    # clause de la nation la plus favorisee vivent en T, U et V. Un controle
    # arrete a la colonne I ne les regarde jamais — et une colonne morte y
    # passerait inapercue exactement comme les ratios de main-d'oeuvre.
    remplies = [c for c in range(4, ws.max_column + 1)
                if ws.cell(row=r, column=c).value is not None]
    cols = [get_column_letter(c) for c in remplies]
    # une hypothese est une ligne dont au moins une cellule de valeur est SAISIE,
    # et dont la valeur saisie est un NOMBRE ou une DATE : une ligne dont la
    # colonne D porte « US$'000 unless stated » ou « FY1 » est un en-tete.
    saisies = [get_column_letter(c) for c in remplies
               if not isinstance(ws.cell(row=r, column=c).value, str)]
    if not saisies:
        continue
    if r in vivantes:
        vivants += 1
        continue
    ns = noms.get((saisies[0], r), [])
    if TOLERES & set(ns):
        tolerees += 1
        continue
    relais = any(cite(c, r, dedans, local=True) for c in cols)
    morts.append((r, lab.strip(), saisies, ns, relais))

print("  classeur   : {}".format(CLASSEUR))
print("  hypotheses consommees : {}".format(vivants))
if tolerees:
    print("  tolerees explicitement: {}".format(tolerees))
print("  HYPOTHESES MORTES     : {}".format(len(morts)))
for r, lab, cols, ns, relais in morts:
    print("      {}{:<4} {:<54} {}{}".format(
        cols[0], r, lab[:54], ",".join(ns) or "(sans nom defini)",
        "  [citee seulement dans la feuille]" if relais else ""))
if morts:
    print("\n  Une hypothese que rien ne consomme oriente le candidat vers une "
          "methode\n  que le modele n'emploie pas. La retirer, ou la brancher.")
sys.exit(1 if morts else 0)

# -*- coding: utf-8 -*-
"""Resout les `#SHARED:N` qu'un generateur de rubric a laisses a la place des formules.

Une formule stockee en formule PARTAGEE dans le .xlsx (`<f t="shared" si="N"/>`)
n'a pas de texte dans la cellule qui la reprend : elle se deduit de la cellule
maitre. Le generateur ne le faisait pas et ecrivait le marqueur brut. Un critere
qui porte `#SHARED:0` ne peut correspondre a aucune formule de candidat.

    python _reparer_shared.py <rubric.json> <golden.xlsx> <sortie.json>
"""
import json, re, sys
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

rubric, golden, sortie = sys.argv[1], sys.argv[2], sys.argv[3]
d = json.load(open(rubric, encoding="utf-8-sig"))
R = d["Rubric"] if "Rubric" in d else d
wf = load_workbook(golden)

resolus, irresolus = 0, []
for sec, crits in R.items():
    for c in crits:
        F = c.get("Formulae") or {}
        cr = c.get("cell_range", "")
        m = re.match(r"^'?([^'!]+)'?!", cr)
        if not m: continue
        f = m.group(1)
        if f not in wf.sheetnames:
            irresolus.append((cr, "feuille absente")); continue
        ws = wf[f]
        for coord, val in list(F.items()):
            if isinstance(val, str) and val.startswith("#SHARED:"):
                reel = ws[coord].value
                if isinstance(reel, str) and reel.startswith("="):
                    F[coord] = reel; resolus += 1
                else:
                    irresolus.append((f"{f}!{coord}", repr(reel)[:40]))

json.dump(d, open(sortie, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"  #SHARED resolus : {resolus}")
if irresolus:
    print(f"  non resolus : {len(irresolus)}")
    for x in irresolus[:8]: print("     ", x)
reste = open(sortie, encoding="utf-8").read().count("#SHARED:")
print(f"  #SHARED restants dans le fichier ecrit : {reste}")

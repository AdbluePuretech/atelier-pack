# -*- coding: utf-8 -*-
"""Les tests des portes, sur le mini-pack synthetique.

    python -m unittest discover -s tests -v          (depuis outils/)

Trois familles :

  - les lecteurs, sur les VRAIS outils lances contre le mini-pack (ceux qui demandent
    Excel ou LibreOffice sont sautes sans eux, et le disent) ;
  - les controles internes, chacun pris en faute sur un fichier abime expres ;
  - la chaine : refus, relais, iterations, peremption, de la porte 1 a juicing.

Quatre controles ne peuvent pas etre verts sur un pack jouet — le niveau (500 logiques
pour un L1), le contrat du prompt et la tracabilite (sept pages du corpus), l'equite
(rubric du corpus). Dans les tests de chaine, et la seulement, ils sont remplaces par
une ligne verte declaree ; leurs lecteurs sont testes a part.
"""
import contextlib
import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

ICI = Path(__file__).resolve().parent
OUTILS = ICI.parent
sys.path.insert(0, str(OUTILS))

import porte  # noqa: E402
import porte_conception as pc  # noqa: E402
import porte_controles as PC  # noqa: E402
import porte_lecteurs as L  # noqa: E402

MINI = ICI / "mini-pack"
GOLD = "build/GoldenSolution - Essai.xlsx"


class AvecCopie(unittest.TestCase):
    """Chaque test travaille sur une copie : le journal et les fichiers abimes n'atteignent jamais la fixture."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="porte-test-"))
        self.pack = self.tmp / "mini-pack"
        shutil.copytree(MINI, self.pack)
        self.M = PC.Manifeste(self.pack / "pack.json")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def p(self, rel):
        return self.pack / rel


def vert(id, libelle):
    return lambda *a, **k: L.juger(id, libelle, 1, 0, "remplace dans les tests de chaine")


# ---- 1. les lecteurs, sur les vrais outils --------------------------------------------

class TestLecteurs(AvecCopie):
    def test_boucles_reelles(self):
        l = L.compter_boucles("t", self.p(GOLD), 1)
        self.assertEqual((l.verdict, l.ecarts), (L.VERT, 0), l.detail)
        self.assertGreater(l.lu, 0)

    def test_boucles_sous_le_plancher(self):
        l = L.compter_boucles("t", self.p(GOLD), 4)
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 3), l.detail)

    def test_hypotheses_mortes(self):
        l = L.inputs_morts("t", self.p(GOLD))
        self.assertEqual(l.verdict, L.VERT, l.detail)
        self.assertIn("4 consommee", l.detail)

    def test_formules_sures(self):
        self.assertEqual(L.formules_sures("t", self.p(GOLD)).verdict, L.VERT)

    def test_niveau_jamais_vert_sur_rien(self):
        l = L.niveau("t", self.p(GOLD), "L1")
        self.assertEqual(l.verdict, L.ROUGE)

    def test_mise_en_page_premier_du_lot(self):
        self.assertEqual(L.mise_en_page("t", self.p(GOLD), None, True).verdict, L.VERT)

    def test_mise_en_page_lot_vide(self):
        vide = self.tmp / "lot_vide"
        vide.mkdir()
        l = L.mise_en_page("t", self.p(GOLD), vide, False)
        self.assertEqual(l.verdict, L.ROUGE)
        self.assertIn("rien lu", l.detail)

    def test_fiche(self):
        l = L.fiche("t", self.p("CONCEPTION.md"))
        self.assertEqual((l.verdict, l.lu), (L.VERT, 11), l.detail)

    def test_fiche_hors_options(self):
        c = self.p("CONCEPTION.md")
        c.write_text(c.read_text(encoding="utf-8").replace("police: Arial", "police: Comic Sans")
                     .replace("bandeau: 1F2A36", "bandeau: FF00FF"), encoding="utf-8")
        l = L.fiche("t", c)
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 2), l.detail)

    def test_golden_100(self):
        l = L.noter("t", self.p(GOLD), self.p("build/rubric.txt"))
        self.assertEqual((l.verdict, l.lu), (L.VERT, 5), l.detail)

    def test_ai_output_pas_100(self):
        l = L.noter("t", self.p("build/AI Output - Essai.xlsx"), self.p("build/rubric.txt"))
        self.assertEqual(l.verdict, L.ROUGE, l.detail)

    def test_contrat_rubric(self):
        self.assertEqual(L.verifier_rubric("t", self.p("build/Rubric - Essai.docx")).verdict, L.VERT)

    def test_contrat_prompt_lit_ses_regles(self):
        l = L.verifier_prompt("t", self.p("build/Prompt - Essai.docx"))
        self.assertGreater(l.lu, 0)

    def test_audit_golden(self):
        self.assertEqual(L.gs_audit("t", self.p(GOLD)).verdict, L.VERT)

    def test_traces_ia(self):
        l = L.traces_ia("t", self.p(GOLD))
        self.assertEqual(l.verdict, L.VERT, l.detail)
        self.assertGreater(l.lu, 0)

    @unittest.skipUnless(L.excel_present(), "Excel absent : les boucles ne se basculent pas")
    def test_amplitude(self):
        l = L.amplitude_boucles("t", self.p(GOLD), self.p("build/rubric.txt"))
        self.assertEqual(l.verdict, L.VERT, l.detail)

    @unittest.skipUnless(L.libreoffice_present(), "LibreOffice absent")
    def test_parite(self):
        l = L.parite_libreoffice("t", self.p(GOLD))
        self.assertEqual(l.verdict, L.VERT, l.detail)


class TestLecteursSortiesIllisibles(unittest.TestCase):
    """Une sortie sans le motif attendu est ROUGE, jamais zero ecart suppose."""

    def test_motif_absent(self):
        with mock.patch.object(L, "outil", return_value=(0, "Traceback: plus rien")):
            for l in (L.compter_boucles("t", "x", 1), L.inputs_morts("t", "x"), L.gs_audit("t", "x"),
                      L.verifier_rubric("t", "x"), L.noter("t", "x", "y"), L.traces_ia("t", "x")):
                self.assertEqual(l.verdict, L.ROUGE, l.id)

    def test_controle_du_pack(self):
        sortie = "blabla\nPORTE equilibre | lu 12 | ecarts 0 | VERT | bilan tenu\nPORTE fx | lu 3 | ecarts 1 | ROUGE | taux\n"
        with mock.patch.object(L, "lancer", return_value=(0, sortie)):
            a, b = L.controle_du_pack("python tester.py", ".", 2)
        self.assertEqual((a.id, a.verdict, a.lu), ("P2.pack.equilibre", L.VERT, 12))
        self.assertEqual(b.verdict, L.ROUGE)

    def test_controle_du_pack_muet(self):
        with mock.patch.object(L, "lancer", return_value=(0, "tout va bien")):
            (l,) = L.controle_du_pack("python tester.py", ".", 2)
        self.assertEqual(l.verdict, L.ROUGE)

    def test_controle_du_pack_qui_ment(self):
        with mock.patch.object(L, "lancer", return_value=(0, "PORTE x | lu 0 | ecarts 0 | VERT |")):
            (l,) = L.controle_du_pack("c", ".", 2)
        self.assertEqual(l.verdict, L.ROUGE, "un VERT declare sans rien lu reste rouge")


# ---- 2. les controles internes, chacun pris en faute ------------------------------------

class TestControlesInternes(AvecCopie):
    def test_cache_plein(self):
        self.assertEqual(PC.cache("t", self.p(GOLD)).verdict, L.VERT)

    def test_cache_jamais_calcule(self):
        import openpyxl
        brut = self.tmp / "brut.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"], wb.active["A2"] = 1, "=A1*2"
        wb.save(brut)
        l = PC.cache("t", brut)
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 1), l.detail)

    def test_iteration_non_armee(self):
        import openpyxl
        brut = self.tmp / "brut.xlsx"
        openpyxl.Workbook().save(brut)
        self.assertEqual(PC.iteration("t", brut).verdict, L.ROUGE)

    def test_ancrages(self):
        ids = ["M1", "M2"]
        self.assertEqual(PC.ancrages_atteints("t", self.p(GOLD), self.p("ancrages.json"), ids).verdict, L.VERT)
        a = json.loads(self.p("ancrages.json").read_text(encoding="utf-8"))
        a[0]["valeur"] = 2.0
        self.p("ancrages.json").write_text(json.dumps(a), encoding="utf-8")
        l = PC.ancrages_atteints("t", self.p(GOLD), self.p("ancrages.json"), ids)
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 1), l.detail)

    def test_input_sheet_qui_cite_le_modele(self):
        import openpyxl
        chemin = self.p("build/InputSheet - Essai.xlsx")
        self.assertEqual(PC.input_sheet("t", chemin).verdict, L.VERT)
        wb = openpyxl.load_workbook(chemin)
        wb.active["F4"] = "=Calc!E4"
        wb.active.row_dimensions[5].outlineLevel = 1
        wb.save(chemin)
        l = PC.input_sheet("t", chemin)
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 2), l.detail)

    def test_certificat_d_une_autre_golden(self):
        cert = self.p("build/audit/certificat.json")
        self.assertEqual(PC.certificat("t", cert, self.p(GOLD)).verdict, L.VERT)
        d = json.loads(cert.read_text(encoding="utf-8"))
        d["empreinte_sha256"], d["couverture_pct"] = "0" * 64, 90
        d["journal"] = [{"cle": "U0", "gravite": "significatif", "constat": "mutation jamais lancee"}]
        cert.write_text(json.dumps(d), encoding="utf-8")
        l = PC.certificat("t", cert, self.p(GOLD))
        self.assertEqual((l.verdict, l.ecarts), (L.ROUGE, 3), l.detail)

    def test_score_sur_une_autre_rubric(self):
        sc, rub = self.p("build/score_ai_output.json"), self.p("build/Rubric - Essai.docx")
        self.assertEqual(PC.score("t", sc, rub, (0, 45), "s").verdict, L.VERT)
        self.assertEqual(PC.score("t", sc, rub, (25, 45), "s").verdict, L.ROUGE)
        self.assertEqual(PC.score("t", sc, self.p("build/Rubric - Essai (juicee).docx"), (0, 45), "s").verdict, L.ROUGE)

    def test_paquet_incomplet(self):
        self.assertEqual(PC.paquet("t", self.M).verdict, L.VERT)
        self.p("build/Prompt - Essai.docx").unlink()
        self.assertEqual(PC.paquet("t", self.M).verdict, L.ROUGE)

    def test_metadonnees(self):
        self.assertEqual(PC.metadonnees("t", self.M).verdict, L.VERT)
        from docx import Document
        chemin = self.p("build/Prompt - Essai.docx")
        d = Document(chemin)
        d.core_properties.author = "Jean Dupont"
        d.save(chemin)
        self.assertEqual(PC.metadonnees("t", self.M).verdict, L.ROUGE)

    def test_note_equite(self):
        rub = self.p("build/Rubric - Essai (juicee).docx")
        self.assertEqual(PC.note_equite("t", self.p("build/note_equite.md"), rub).verdict, L.VERT)
        self.p("build/note_equite.md").write_text("# Note\n\n- une seule entree\n", encoding="utf-8")
        self.assertEqual(PC.note_equite("t", self.p("build/note_equite.md"), rub).verdict, L.ROUGE)


class TestPorte1(AvecCopie):
    def lignes(self):
        return {l.id: l for l in PC.porte_1(self.M, None)}

    def test_conception_saine(self):
        rouges = [l for l in self.lignes().values() if l.verdict != L.VERT]
        self.assertEqual(rouges, [])

    def test_titre_manquant(self):
        c = self.p("CONCEPTION.md")
        c.write_text(c.read_text(encoding="utf-8").replace("## Frontieres", "## Autre chose"), encoding="utf-8")
        self.assertEqual(self.lignes()["P1.plan"].verdict, L.ROUGE)

    def test_citation_inventee(self):
        c = self.p("CONCEPTION.md")
        c.write_text(c.read_text(encoding="utf-8").replace("interest accrues on those net proceeds",
                                                           "a tax election on the sale"), encoding="utf-8")
        self.assertEqual(self.lignes()["P1.graine"].verdict, L.ROUGE)

    def test_plancher_de_boucles_L2(self):
        d = json.loads(self.p("pack.json").read_text(encoding="utf-8"))
        d["niveau"] = "L2"
        self.p("pack.json").write_text(json.dumps(d), encoding="utf-8")
        self.M = PC.Manifeste(self.p("pack.json"))
        l = self.lignes()
        self.assertEqual(l["P1.boucles"].verdict, L.ROUGE)
        self.assertEqual(l["P1.niveau"].verdict, L.ROUGE)

    def test_budget_au_dessus_du_plafond(self):
        c = self.p("CONCEPTION.md")
        c.write_text(c.read_text(encoding="utf-8").replace("transcription : 15 %", "transcription : 37 %"), encoding="utf-8")
        self.assertEqual(self.lignes()["P1.budget"].verdict, L.ROUGE)

    def test_ancrage_incomplet(self):
        a = json.loads(self.p("ancrages.json").read_text(encoding="utf-8"))
        a[2]["nature"], a[2]["pourquoi"] = "X", ""
        self.p("ancrages.json").write_text(json.dumps(a), encoding="utf-8")
        self.assertEqual(self.lignes()["P1.ancrages"].ecarts, 2)

    def test_livrable_absent(self):
        self.p("graine.md").unlink()
        (l,) = PC.porte_1(self.M, None)
        self.assertEqual(l.verdict, L.NON_FAIT)


# ---- 3. la chaine ----------------------------------------------------------------------

REMPLACES = dict(niveau=vert("P2.niveau", "niveau tenu"), verifier_prompt=vert("P4.prompt", "contrat du prompt"),
                 tracabilite=vert("P.tracabilite", "tracabilite"), dossier_equite=vert("P5.equite", "equite"),
                 amplitude_boucles=vert("P4.amplitude", "chaque boucle notee"),
                 parite_libreoffice=vert("P5.parite", "parite LibreOffice"), excel_present=lambda: True)


class TestChaine(AvecCopie):
    def setUp(self):
        super().setUp()
        self.patches = [mock.patch.object(L, k, v) for k, v in REMPLACES.items()]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()
        super().tearDown()

    def porte(self, nom, iteration=None):
        with contextlib.redirect_stdout(io.StringIO()) as out:
            code = porte.lancer_porte(PC.Manifeste(self.p("pack.json")), nom, iteration)
        return code, out.getvalue()

    def relais(self, nom, decision="valide", iteration=None):
        with contextlib.redirect_stdout(io.StringIO()) as out:
            code = porte.relais(PC.Manifeste(self.p("pack.json")), nom, iteration, decision, "test")
        return code, out.getvalue()

    def jusqua_porte_2(self):
        for nom, it in (("1", None), ("2", 0), ("2", 1), ("2", 2)):
            self.assertEqual(self.porte(nom, it)[0], 0, (nom, it))
            self.assertEqual(self.relais(nom, iteration=it)[0], 0)

    def test_refus_sans_porte_precedente(self):
        code, out = self.porte("2")
        self.assertEqual(code, 2)
        self.assertIn("porte 1 jamais passee", out)

    def test_refus_sans_relais(self):
        self.porte("1")
        code, out = self.porte("2", 0)
        self.assertEqual(code, 2)
        self.assertIn("sans relais valide", out)

    def test_iteration_sautee(self):
        self.porte("1")
        self.relais("1")
        code, out = self.porte("2", 1)
        self.assertEqual(code, 2)
        self.assertIn("iteration 0 jamais passee", out)

    def test_relais_sur_porte_fermee(self):
        self.p("graine.md").unlink()
        self.assertEqual(self.porte("1")[0], 1)
        code, out = self.relais("1")
        self.assertEqual(code, 2)
        self.assertIn("n'est pas ouverte", out)

    def test_chaine_complete(self):
        self.jusqua_porte_2()
        for nom in ("2", "3", "4", "5", "juicing"):
            code, out = self.porte(nom)
            self.assertEqual(code, 0, out)
            self.assertEqual(self.relais(nom)[0], 0)
        copies = sorted(x.name for x in self.p("build/iterations").iterdir())
        self.assertEqual(copies, [f"GS it0{k} - Essai.xlsx" for k in range(3)])

    def test_golden_modifiee_rend_perimees_les_suivantes(self):
        self.jusqua_porte_2()
        for nom in ("2", "3"):
            self.porte(nom)
            self.relais(nom)
        gold = self.p(GOLD)
        with zipfile.ZipFile(gold, "a") as z:
            z.writestr("customXml/retouche.xml", "<x/>")
        etat = {(l["porte"], l["iteration"]): l for l in porte.Journal(PC.Manifeste(self.p("pack.json"))).etat()}
        self.assertFalse(etat[("2", 2)]["perimee"], "une iteration suit sa copie figee, pas la golden vivante")
        self.assertTrue(etat[("2", None)]["perimee"])
        self.assertTrue(etat[("3", None)]["perimee"])
        code, out = self.porte("4")
        self.assertEqual(code, 2)
        self.assertIn("perimee", out)

    def test_cellules_d_ancrage_ne_perimant_pas_la_porte_1(self):
        self.porte("1")
        self.relais("1")
        a = json.loads(self.p("ancrages.json").read_text(encoding="utf-8"))
        a[0]["cellule"] = "Calc!F5"
        self.p("ancrages.json").write_text(json.dumps(a), encoding="utf-8")
        self.assertEqual(self.porte("2", 0)[0], 0, "remplir une cellule d'ancrage est le travail de la phase 2")
        a[0]["valeur"] = 3.0
        self.p("ancrages.json").write_text(json.dumps(a), encoding="utf-8")
        code, out = self.porte("2", 1)
        self.assertEqual(code, 2)
        self.assertIn("perimee", out)

    def test_modifie_exige_les_decisions(self):
        self.porte("1")
        self.relais("1")
        self.porte("2", 0)
        self.relais("2", "modifie", 0)
        code, out = self.porte("2", 0)
        self.assertEqual(code, 1)
        self.assertIn("it00.json est absent", out)


class TestInitEtManifeste(unittest.TestCase):
    def test_init_n_ecrase_rien(self):
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "CONCEPTION.md").write_text("la mienne", encoding="utf-8")
            with contextlib.redirect_stdout(io.StringIO()):
                porte.init(d)
            self.assertEqual((Path(d) / "CONCEPTION.md").read_text(encoding="utf-8"), "la mienne")
            self.assertTrue((Path(d) / "pack.json").exists())
            with self.assertRaises(PC.ManifesteInvalide):
                PC.Manifeste(Path(d) / "pack.json")

    def test_plan_vide_a_les_dix_titres(self):
        with tempfile.TemporaryDirectory() as d:
            c = Path(d) / "c.md"
            c.write_text(pc.PLAN_VIDE.format(nom="X"), encoding="utf-8")
            self.assertEqual(len(pc.Conception(c).titres_presents()), 10)


if __name__ == "__main__":
    unittest.main()

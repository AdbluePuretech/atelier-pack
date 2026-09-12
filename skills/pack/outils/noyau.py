# -*- coding: utf-8 -*-
"""Noyau du moteur d'audit : lecture du classeur, analyse des references,
graphe de dependances, composantes fortement connexes, familles de formules.

Aucun controle ici, seulement la matiere premiere que les autres modules
consomment. Rien ne fait confiance au generateur : tout se lit dans le
fichier livre.
"""
import re
import datetime
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

_CHAINE = r'"(?:[^"]|"")*"'
_FEUILLE = r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)"
_CELL = r"\$?[A-Z]{1,3}\$?[0-9]+"
_NOM = r"[A-Za-z_][A-Za-z0-9_.]*"

JETON = re.compile(
    r"(?P<chaine>" + _CHAINE + r")"
    r"|(?P<qplage>" + _FEUILLE + r"!" + _CELL + r":" + _CELL + r")"
    r"|(?P<qcell>" + _FEUILLE + r"!" + _CELL + r")"
    r"|(?P<plage>(?<![A-Z0-9_.!])" + _CELL + r":" + _CELL + r")"
    r"|(?P<cell>(?<![A-Z0-9_.!])" + _CELL + r"(?![A-Z0-9_(]))"
    r"|(?P<appel>" + _NOM + r"\s*\()"
    r"|(?P<nom>(?<![A-Z0-9_.!])" + _NOM + r"(?!\s*\())"
)

CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?([0-9]+)")


def _pos(ref):
    m = CELL_RE.fullmatch(ref)
    return int(m.group(2)), column_index_from_string(m.group(1))


def _defeuille(s):
    s = s.strip()
    if s.startswith("'") and s.endswith("'"):
        s = s[1:-1].replace("''", "'")
    return s


class Classeur:
    """Le classeur livre, lu deux fois : formules et valeurs en cache."""

    def __init__(self, chemin, plafond_plage=40000):
        self.chemin = chemin
        self.plafond_plage = plafond_plage
        self.wf = openpyxl.load_workbook(chemin, read_only=False, data_only=False)
        self.wv = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
        self.feuilles = [ws.title for ws in self.wf.worksheets]

        self.formules = {}
        self.durs = {}
        self.textes = {}
        self.dates = {}
        for ws in self.wf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None:
                        continue
                    cle = (ws.title, c.row, c.column)
                    if isinstance(v, str) and v.startswith("="):
                        self.formules[cle] = v
                    elif isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
                        self.dates[cle] = v
                    elif isinstance(v, (int, float)):
                        self.durs[cle] = v
                    elif isinstance(v, str):
                        self.textes[cle] = v

        self.valeurs = {}
        for ws in self.wv.worksheets:
            for row in ws.iter_rows(values_only=False):
                for c in row:
                    if c.value is not None:
                        self.valeurs[(ws.title, c.row, c.column)] = c.value

        self.noms = self._resoudre_noms()

    def _resoudre_noms(self):
        out = {}
        try:
            items = list(self.wf.defined_names.items())
        except AttributeError:
            items = [(d.name, d) for d in self.wf.defined_names.definedName]
        for nom, dn in items:
            cibles = []
            try:
                for feuille, ref in dn.destinations:
                    cibles.extend(self._etendre(feuille, ref))
            except Exception:
                pass
            out[nom] = cibles
        return out

    def _etendre(self, feuille, ref):
        ref = ref.replace("$", "")
        cases = []
        for morceau in ref.split(","):
            if ":" in morceau:
                a, b = morceau.split(":", 1)
                try:
                    l1, c1 = _pos(a)
                    l2, c2 = _pos(b)
                except Exception:
                    continue
                if (abs(l2 - l1) + 1) * (abs(c2 - c1) + 1) > self.plafond_plage:
                    continue
                for l in range(min(l1, l2), max(l1, l2) + 1):
                    for c in range(min(c1, c2), max(c1, c2) + 1):
                        cases.append((feuille, l, c))
            else:
                try:
                    l, c = _pos(morceau)
                except Exception:
                    continue
                cases.append((feuille, l, c))
        return cases

    def references(self, cle, formule):
        feuille_courante = cle[0]
        out = []
        for m in JETON.finditer(formule):
            k = m.lastgroup
            if k in ("chaine", "appel"):
                continue
            t = m.group()
            if k in ("qplage", "qcell"):
                f, r = t.split("!", 1)
                out.extend(self._etendre(_defeuille(f), r))
            elif k in ("plage", "cell"):
                out.extend(self._etendre(feuille_courante, t))
            elif k == "nom":
                out.extend(self.noms.get(t, []))
        return out

    def graphe(self):
        """Arcs cellule -> precedents, restreints aux cellules de formule.

        Une constante est un puits : elle n'appartient a aucun cycle. En
        l'ecartant, le graphe fond et les composantes deviennent calculables.
        """
        idx = {cle: i for i, cle in enumerate(self.formules)}
        adj = [[] for _ in idx]
        for cle, f in self.formules.items():
            i = idx[cle]
            vus = set()
            for r in self.references(cle, f):
                j = idx.get(r)
                if j is not None and j != i and j not in vus:
                    vus.add(j)
                    adj[i].append(j)
        self.idx = idx
        self.rev_idx = {i: c for c, i in idx.items()}
        self.adj = adj
        return adj

    def composantes(self, adj):
        """Tarjan iteratif. Retourne les composantes fortement connexes."""
        n = len(adj)
        index = [-1] * n
        bas = [0] * n
        surpile = bytearray(n)
        pile = []
        out = []
        compteur = 0
        for depart in range(n):
            if index[depart] != -1:
                continue
            travail = [(depart, 0)]
            while travail:
                v, pi = travail[-1]
                if pi == 0:
                    index[v] = bas[v] = compteur
                    compteur += 1
                    pile.append(v)
                    surpile[v] = 1
                recurse = False
                av = adj[v]
                for i in range(pi, len(av)):
                    w = av[i]
                    if index[w] == -1:
                        travail[-1] = (v, i + 1)
                        travail.append((w, 0))
                        recurse = True
                        break
                    elif surpile[w]:
                        if index[w] < bas[v]:
                            bas[v] = index[w]
                if recurse:
                    continue
                if bas[v] == index[v]:
                    comp = []
                    while True:
                        w = pile.pop()
                        surpile[w] = 0
                        comp.append(w)
                        if w == v:
                            break
                    out.append(comp)
                travail.pop()
                if travail:
                    u = travail[-1][0]
                    if bas[v] < bas[u]:
                        bas[u] = bas[v]
        return out


_SK_REF = re.compile(_FEUILLE + r"!" + _CELL + r"(?::" + _CELL + r")?")
_SK_LOC = re.compile(r"(?<![A-Z0-9_.!])" + _CELL + r"(?::" + _CELL + r")?(?![A-Z0-9_(])")
_SK_NOM = re.compile(r"\b[a-z]{1,3}_[A-Za-z0-9_]+\b")
_SK_NUM = re.compile(r"(?<![A-Za-z0-9_.])[0-9]+(?:\.[0-9]+)?(?:[eE][-+]?[0-9]+)?")


def squelette(f):
    """Forme structurelle : references -> X, plages nommees -> N, litteraux -> #."""
    s = _SK_REF.sub("X", f)
    s = _SK_LOC.sub("X", s)
    s = _SK_NOM.sub("N", s)
    s = _SK_NUM.sub("#", s)
    s = re.sub(r"X(?:\s*[:,]\s*X)+", "X", s)
    return re.sub(r"\s+", "", s)


def libelle(cle):
    return f"{cle[0]}!{get_column_letter(cle[2])}{cle[1]}"


_TOTAL = re.compile(
    r"(sum of every|grand total|total .*(violation|residual|check)|"
    r"somme de tous|total general)", re.I)
_SOMME = re.compile(r"\bSUM\s*\(", re.I)


def trouver_batterie(C):
    """L'adresse du total de la batterie de controles, trouvee toute seule.

    Trois modules la demandaient en parametre, et rien ne disait comment
    l'obtenir : l'utilisateur devait vider une feuille a la main pour lancer
    une etape. Un outil qui exige une information qu'il pourrait deduire n'est
    pas fini.

    On cherche, sur une feuille de controles, la ligne dont le libelle annonce
    un total general, et on prend sa premiere cellule de formule. A defaut, la
    derniere formule de la feuille qui somme une plage - un total est presque
    toujours en bas.
    """
    feuilles = [f for f in C.feuilles
                if "check" in f.lower() or "controle" in f.lower()]
    for feuille in feuilles or C.feuilles:
        candidats = []
        for (f, l, c), v in C.textes.items():
            if f != feuille or not isinstance(v, str):
                continue
            if _TOTAL.search(v):
                for cc in range(c + 1, c + 16):
                    if (f, l, cc) in C.formules:
                        candidats.append((l, f"'{f}'!{get_column_letter(cc)}{l}"))
                        break
        if candidats:
            return max(candidats)[1]        # le plus bas de la feuille
    for feuille in feuilles:
        sommes = [(l, c) for (f, l, c) in C.formules
                  if f == feuille and _SOMME.search(C.formules[(f, l, c)])]
        if sommes:
            l, c = max(sommes)
            return f"'{feuille}'!{get_column_letter(c)}{l}"
    return None

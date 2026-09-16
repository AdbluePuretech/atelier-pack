# -*- coding: utf-8 -*-
"""Le registre des portes : pour chaque porte, la liste de ses controles.

Deux sortes de controles vivent ici :

  - ceux qui lancent un outil de la skill, par un lecteur de `porte_lecteurs.py` ;
  - ceux qui n'ont pas d'outil et se lisent directement dans les fichiers : le cache
    de la golden, l'iteration armee, les ancrages atteints, l'input sheet, le
    certificat, la feuille de score, le paquet, les metadonnees, les decisions
    d'atelier.

Un controle qui ne peut pas tourner (livrable absent, champ vide, logiciel absent)
rend NON FAIT : la porte reste fermee et dit pourquoi, elle ne suppose rien.
"""
import hashlib
import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import porte_conception as pc
import porte_lecteurs as L
from porte_lecteurs import juger, non_fait

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
COUVERTURE_MIN = 95
BANDE_JUICING = (20, 40)


def empreinte(chemin):
    h = hashlib.sha256()
    with open(chemin, "rb") as fh:
        for bloc in iter(lambda: fh.read(1 << 20), b""):
            h.update(bloc)
    return h.hexdigest()


class ManifesteInvalide(Exception):
    pass


class Manifeste:
    def __init__(self, chemin):
        self.chemin = Path(chemin).resolve()
        self.dossier = self.chemin.parent
        try:
            self.d = json.loads(self.chemin.read_text(encoding="utf-8"))
        except FileNotFoundError:
            raise ManifesteInvalide(f"manifeste introuvable : {self.chemin}")
        except json.JSONDecodeError as e:
            raise ManifesteInvalide(f"pack.json n'est pas du JSON valide : {e}")
        if not self.rempli(self.d.get("nom")):
            raise ManifesteInvalide("champ « nom » absent ou non rempli")
        if self.d.get("niveau") not in pc.PLANCHER_BOUCLES:
            raise ManifesteInvalide(f"champ « niveau » : {self.d.get('niveau')!r}, attendu L1, L2 ou L3")

    @staticmethod
    def rempli(v):
        return isinstance(v, str) and v.strip() and not re.search(r"<[^>]*>", v)

    @property
    def nom(self):
        return self.d["nom"]

    @property
    def niveau(self):
        return self.d["niveau"]

    @property
    def premier_du_lot(self):
        return bool(self.d.get("premier_du_lot"))

    def chemin_de(self, cle, section=None):
        v = (self.d.get(section) or {}).get(cle) if section else self.d.get(cle)
        if not self.rempli(v):
            return None
        p = Path(v)
        return p if p.is_absolute() else (self.dossier / p)

    def exiger(self, id, libelle, cles, section=None):
        """Rend (chemins, None) ou (None, ligne NON FAIT)."""
        chemins = []
        for cle in cles:
            nom = f"{section}.{cle}" if section else cle
            p = self.chemin_de(cle, section)
            if p is None:
                return None, non_fait(id, libelle, f"champ « {nom} » absent ou non rempli dans pack.json")
            if not p.exists():
                return None, non_fait(id, libelle, f"fichier introuvable pour « {nom} » : {p}")
            chemins.append(p)
        return chemins, None

    def lot(self, id, libelle):
        if self.premier_du_lot:
            return None, None
        p = self.chemin_de("lot")
        if p is None:
            return None, non_fait(id, libelle, "champ « lot » absent, et « premier_du_lot » n'est pas declare")
        if not p.is_dir():
            return None, non_fait(id, libelle, f"dossier du lot introuvable : {p}")
        return p, None


# ---- controles internes -------------------------------------------------------

def feuilles_xml(chemin):
    with zipfile.ZipFile(chemin) as z:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        cible = {r.get("Id"): r.get("Target") for r in rels}
        for s in wb.iter(NS + "sheet"):
            rid = s.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            t = cible.get(rid, "")
            t = t.lstrip("/") if t.startswith("/") else "xl/" + t
            if t in z.namelist():
                yield s.get("name"), z.read(t)


def cache(id, golden):
    """0 erreur en cache, 0 formule sans valeur en cache — lu dans le XML, sans Excel.

    openpyxl ecrit une formule avec une valeur vide et sans type ; Excel ecrit une chaine
    vide avec t="str". Le premier cas est un cache jamais calcule, le second un resultat.
    """
    lues = formules = erreurs = sans = 0
    exemples = []
    for nom, xml in feuilles_xml(golden):
        for c in ET.fromstring(xml).iter(NS + "c"):
            if len(c):
                lues += 1
            if c.find(NS + "f") is None:
                continue
            formules += 1
            t, v = c.get("t"), c.find(NS + "v")
            if t == "e":
                erreurs += 1
                if len(exemples) < 3:
                    exemples.append(f"{nom}!{c.get('r')} {v.text if v is not None else ''}")
            elif v is None or (v.text is None and t not in ("str", "inlineStr")):
                sans += 1
                if len(exemples) < 3:
                    exemples.append(f"{nom}!{c.get('r')} sans cache")
    return juger(id, "cache de la golden", lues, erreurs + sans,
                 f"{erreurs} erreur(s), {sans} formule(s) sans cache sur {formules}"
                 + (f" ; {', '.join(exemples)}" if exemples else ""), [golden])


def iteration(id, classeur, libelle="calcul iteratif arme"):
    with zipfile.ZipFile(classeur) as z:
        xml = z.read("xl/workbook.xml").decode("utf-8", errors="replace")
    arme = bool(re.search(r"<calcPr\b[^>]*\biterate=\"(1|true)\"", xml))
    return juger(id, libelle, 1, 0 if arme else 1, "iterate arme dans calcPr" if arme else "calcPr sans iterate", [classeur])


def adresse(cellule):
    f, cel = str(cellule).rsplit("!", 1)
    return f.strip("'"), cel.replace("$", "")


def ancrages_atteints(id, golden, ancrages_json, ids, libelle="ancrages atteints"):
    from openpyxl import load_workbook
    ancrages = [a for a in pc.lire_ancrages(ancrages_json) if a.get("mecanique") in ids]
    if not ancrages:
        return juger(id, libelle, 0, 0, f"aucun ancrage pour {', '.join(ids) or 'aucune mecanique'}", [golden, ancrages_json])
    wb = load_workbook(golden, data_only=True)
    rates = []
    for a in ancrages:
        try:
            f, cel = adresse(a["cellule"])
            v = wb[f][cel].value
        except Exception:
            rates.append(f"{a.get('id')} : cellule {a.get('cellule')!r} illisible")
            continue
        if not pc.dans_tolerance(v, a):
            rates.append(f"{a.get('id')} {a['cellule']} = {v}, attendu {a['valeur']}")
    return juger(id, libelle, len(ancrages), len(rates),
                 f"{len(ancrages) - len(rates)} / {len(ancrages)} atteints" + (" ; " + "; ".join(rates[:4]) if rates else ""),
                 [golden, ancrages_json])


REF_FEUILLE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][\w.]*))!")


def input_sheet(id, chemin):
    from openpyxl import load_workbook
    wb = load_workbook(chemin)
    feuilles = set(wb.sheetnames)
    lues, hors, externes, groupes = 0, [], 0, 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                lues += 1
                if isinstance(c.value, str) and c.value.startswith("="):
                    if re.search(r"\[\d+\]", c.value):
                        externes += 1
                    for m in REF_FEUILLE.finditer(c.value):
                        nom = (m.group(1) or m.group(2)).replace("''", "'")
                        if nom not in feuilles:
                            hors.append(f"{ws.title}!{c.coordinate} -> {nom}")
        groupes += sum(1 for d in ws.row_dimensions.values() if (d.outlineLevel or 0) > 0)
        groupes += sum(1 for d in ws.column_dimensions.values() if (d.outlineLevel or 0) > 0)
    return juger(id, "input sheet extraite", lues, len(hors) + externes + groupes,
                 f"{len(hors)} formule(s) vers le modele, {externes} lien(s) externe(s), {groupes} groupement(s), "
                 f"{len(feuilles)} onglet(s)" + (f" ; {', '.join(hors[:3])}" if hors else ""), [chemin])


def certificat(id, cert_json, golden):
    cert = json.loads(Path(cert_json).read_text(encoding="utf-8"))
    ecarts, notes = 0, []
    if cert.get("empreinte_sha256") != empreinte(golden):
        ecarts += 1
        notes.append("le certificat ne porte pas l'empreinte de la golden actuelle")
    couv = cert.get("couverture_pct") or 0
    if couv < COUVERTURE_MIN:
        ecarts += 1
        notes.append(f"couverture {couv} % < {COUVERTURE_MIN} %")
    journal = cert.get("journal") or []
    non_faits = [e["cle"] for e in journal if str(e.get("cle", "")).endswith("0")]
    critiques = [e["cle"] for e in journal if e.get("gravite") == "critique" and e["cle"] not in non_faits]
    ecarts += len(non_faits) + len(critiques)
    if non_faits:
        notes.append("non fait : " + ", ".join(non_faits))
    if critiques:
        notes.append("defaut critique : " + ", ".join(critiques))
    return juger(id, "certificat de la golden", cert.get("cellules_de_calcul") or 0, ecarts,
                 f"couverture {couv} %" + (" ; " + "; ".join(notes) if notes else ""), [cert_json, golden])


def score(id, score_json, rubric, bande, libelle):
    s = json.loads(Path(score_json).read_text(encoding="utf-8"))
    pct, notes, ecarts = s.get("pourcentage"), [], 0
    if not isinstance(pct, (int, float)):
        return juger(id, libelle, None, 0, "champ « pourcentage » absent", [score_json, rubric])
    if not bande[0] <= pct <= bande[1]:
        ecarts += 1
        notes.append(f"{pct} % hors de la bande {bande[0]}-{bande[1]} %")
    if s.get("rubric_sha256") != empreinte(rubric):
        ecarts += 1
        notes.append("score rendu sur une autre version de la rubric")
    return juger(id, libelle, int(s.get("pool") or 0), ecarts,
                 f"{pct} % ({s.get('source', '?')}, {s.get('date', '?')})" + (" ; " + "; ".join(notes) if notes else ""),
                 [score_json, rubric])


def paquet(id, M):
    liste = M.d.get("paquet") or []
    if not liste:
        return non_fait(id, "paquet complet", "champ « paquet » vide dans pack.json")
    chemins = [(Path(x) if Path(x).is_absolute() else M.dossier / x) for x in liste]
    manquants = [p.name for p in chemins if not p.exists()]
    return juger(id, "paquet complet", len(chemins), len(manquants),
                 "manquants : " + ", ".join(manquants) if manquants else f"{len(chemins)} fichiers presents",
                 [p for p in chemins if p.exists()])


TRACES = re.compile(r"[A-Za-z]:\\Users\\|OneDrive|sharepoint|/Users/", re.I)


def metadonnees(id, M):
    liste = [(Path(x) if Path(x).is_absolute() else M.dossier / x) for x in (M.d.get("paquet") or [])]
    lus, fautes = 0, []
    for p in liste:
        if p.suffix.lower() not in (".xlsx", ".docx") or not p.exists():
            continue
        lus += 1
        with zipfile.ZipFile(p) as z:
            for part in [n for n in z.namelist() if n.startswith("docProps/")] + ["xl/workbook.xml"]:
                if part not in z.namelist():
                    continue
                xml = z.read(part).decode("utf-8", errors="replace")
                for balise in ("dc:creator", "cp:lastModifiedBy", "Company", "Manager"):
                    m = re.search(rf"<{balise}>([^<]+)</{balise}>", xml)
                    if m and m.group(1).strip():
                        fautes.append(f"{p.name} {balise} = {m.group(1)[:30]}")
                if TRACES.search(xml):
                    fautes.append(f"{p.name} {part} : chemin machine ou compte")
    return juger(id, "metadonnees nettoyees", lus, len(fautes), "; ".join(fautes[:4]) if fautes else f"{lus} fichiers propres",
                 [p for p in liste if p.exists()])


def egaux(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) <= 1e-6 * max(1, abs(a), abs(b))
    return str(a) == str(b)


def atelier(id, fichier, golden):
    from openpyxl import load_workbook
    d = json.loads(Path(fichier).read_text(encoding="utf-8"))
    modifs = [m for m in d.get("modifications", []) if m.get("cellule")]
    sans, ecart = [], []
    wb = None
    for m in modifs:
        dec = m.get("decision")
        if dec not in ("reportee", "abandonnee"):
            sans.append(f"{m['feuille']}!{m['cellule']}")
            continue
        if dec == "reportee":
            wb = wb or load_workbook(golden)
            v = wb[m["feuille"]][m["cellule"]].value if m["feuille"] in wb.sheetnames else None
            if not egaux(v, m.get("apres")):
                ecart.append(f"{m['feuille']}!{m['cellule']} = {v!r}, reporte {m.get('apres')!r}")
    return juger(id, "decisions d'atelier", len(modifs), len(sans) + len(ecart),
                 (f"{len(sans)} sans decision" + (f" ({', '.join(sans[:3])})" if sans else "")
                  + f", {len(ecart)} reportee(s) absente(s) de la golden" + (f" ({'; '.join(ecart[:3])})" if ecart else "")),
                 [fichier, golden])


def note_equite(id, note, rubric):
    from docx import Document
    textes = [p.text for p in Document(rubric).paragraphs]
    gates = sum(1 for t in textes if re.search(r"\bgate\b", t, re.I) and re.match(r"^\s*\[\+", t))
    keystones = sum(1 for t in textes if re.match(r"^\s*\[\+4\]", t))
    attendu = gates + keystones
    entrees = sum(1 for l in Path(note).read_text(encoding="utf-8").splitlines()
                  if re.match(r"^\s*([-*]\s+\S|\|\s*[^|\-\s])", l))
    return juger(id, "note d'equite", attendu, max(0, attendu - entrees),
                 f"{entrees} entree(s) pour {gates} gate(s) et {keystones} keystone(s)", [note, rubric])


# ---- les portes ------------------------------------------------------------------

def _controles_du_pack(M, porte):
    lignes = []
    for commande in (M.d.get("controles_du_pack") or {}).get(str(porte), []):
        lignes += L.controle_du_pack(commande, M.dossier, porte)
    return lignes


def _sur(id, libelle, fn, fichiers):
    """Un controle qui plante rend ROUGE avec son erreur, jamais une exception qui tue la porte."""
    try:
        return fn()
    except Exception as e:
        return juger(id, libelle, None, 0, f"{type(e).__name__}: {e}", fichiers)


def porte_1(M, J):
    out = []
    chemins, nf = M.exiger("P1.plan", "plan de la conception", ["conception", "graine", "ancrages"])
    if nf:
        return [nf]
    conc, graine, anc = chemins
    C = pc.Conception(conc)
    presents = C.titres_presents()
    manquants = [t for t in pc.TITRES if t not in presents]
    out.append(juger("P1.plan", "plan de la conception", len(pc.TITRES), len(manquants),
                     "titres manquants : " + ", ".join(manquants) if manquants else "10 titres presents", [conc]))
    nv = C.niveau()
    out.append(juger("P1.niveau", "niveau vise", 1, 0 if nv == M.niveau else 1,
                     f"conception {nv or 'absent'}, manifeste {M.niveau}", [conc]))
    texte_graine = Path(graine).read_text(encoding="utf-8")
    rangs = C.graine()
    absentes = [pc.colonne(r, "Element") or "?" for r in rangs
                if not pc.citation_trouvee(pc.colonne(r, "Passage de la graine"), texte_graine)]
    out.append(juger("P1.graine", "graine -> pack", len(rangs), len(absentes),
                     "passage introuvable dans la graine : " + ", ".join(absentes[:5]) if absentes
                     else f"{len(rangs)} elements rattaches a la graine", [conc, graine]))
    meca = C.mecaniques()
    ids = C.ids()
    sans_piege = [pc.colonne(r, "Id") or "?" for r in meca if not pc.colonne(r, "Piege nomme")]
    mal = len(meca) - len(ids) + (len(ids) - len(set(ids)))
    out.append(juger("P1.pieges", "un piege par mecanique", len(meca), len(sans_piege) + mal,
                     f"{len(ids)} mecanique(s), sans piege : {', '.join(sans_piege) or 'aucune'}, id mal forme ou double : {mal}", [conc]))
    boucles = C.boucles()
    valides = [b for b in boucles if pc.colonne(b, "Convention des pieces") and pc.colonne(b, "Mecanique") in ids]
    plancher = pc.PLANCHER_BOUCLES[M.niveau]
    out.append(juger("P1.boucles", "boucles nommees", max(len(boucles), 1 if boucles else 0),
                     (len(boucles) - len(valides)) + max(0, plancher - len(valides)),
                     f"{len(valides)} boucle(s) complete(s), plancher {M.niveau} : {plancher}", [conc]))

    def _ancrages():
        A = pc.lire_ancrages(anc)
        defauts = [d for a in A for d in pc.defauts_ancrage(a, ids, exige_cellule=False)]
        sans = [i for i in ids if not any(a.get("mecanique") == i for a in A)]
        return juger("P1.ancrages", "ancrages", len(A), len(defauts) + len(sans),
                     "; ".join((defauts + [f"{i} sans ancrage" for i in sans])[:5]) or f"{len(A)} ancrages complets", [anc])
    out.append(_sur("P1.ancrages", "ancrages", _ancrages, [anc]))
    b = C.budget()
    out.append(juger("P1.budget", "budget de discrimination", 1 if b is not None else 0,
                     0 if b is not None and b <= pc.BUDGET_MAX else 1,
                     f"transcription {b} %, plafond {pc.BUDGET_MAX} %" if b is not None else "ligne du budget absente", [conc]))
    out.append(L.fiche("P1.fiche", conc))
    lot, nf = M.lot("P1.distance", "distance de la fiche au lot")
    out.append(nf or L.distance_fiche("P1.distance", conc, lot, M.premier_du_lot))
    return out + _controles_du_pack(M, 1)


def _ligne(M, id, libelle, cles, fn):
    """Une ligne qui a besoin de livrables : NON FAIT pour elle seule s'ils manquent, la porte continue."""
    chemins, nf = M.exiger(id, libelle, cles)
    if nf:
        return nf
    return _sur(id, libelle, lambda: fn(*chemins), chemins)


def porte_2(M, J, iteration=None, iterations_validees=None):
    chemins, nf = M.exiger("P2.cache", "cache de la golden", ["golden", "conception"])
    if nf:
        return [nf]
    gold, conc = chemins
    C = pc.Conception(conc)
    ids = C.ids()
    p = "P2i" if iteration is not None else "P2"
    out = [_sur(f"{p}.cache", "cache de la golden", lambda: cache(f"{p}.cache", gold), [gold]),
           _sur(f"{p}.iteration", "calcul iteratif arme", lambda: iteration_(f"{p}.iteration", gold), [gold])]
    if iteration is not None:
        sous = ids[:iteration]
        if iteration > 0:
            out.append(_ligne(M, f"{p}.ancrages", "ancrages atteints", ["ancrages"],
                              lambda anc: ancrages_atteints(f"{p}.ancrages", gold, anc, sous)))
            declarees = sum(1 for b in C.boucles() if pc.colonne(b, "Mecanique") in sous)
            if declarees:
                out.append(L.compter_boucles(f"{p}.boucles", gold, declarees, f" (boucles declarees pour {', '.join(sous)})"))
        fichier = M.dossier / "build" / "ateliers" / f"it{iteration:02d}.json"
        if J.relais_le_plus_recent("2", iteration) == "modifie":
            out.append(_sur(f"{p}.atelier", "decisions d'atelier", lambda: atelier(f"{p}.atelier", fichier, gold), [fichier])
                       if fichier.exists() else non_fait(f"{p}.atelier", "decisions d'atelier",
                                                         f"le owner a modifie, mais {fichier.name} est absent"))
        return out + _controles_du_pack(M, "2i")
    out.append(_ligne(M, "P2.ancrages", "ancrages atteints", ["ancrages"],
                      lambda anc: ancrages_atteints("P2.ancrages", gold, anc, ids)))
    plancher = pc.PLANCHER_BOUCLES[M.niveau]
    out.append(L.compter_boucles("P2.boucles", gold, max(plancher, len(C.boucles())), f" (plancher {plancher}, declarees {len(C.boucles())})"))
    out.append(L.inputs_morts("P2.inputs_morts", gold))
    out.append(L.formules_sures("P2.formules_sures", gold))
    out.append(L.niveau("P2.niveau", gold, M.niveau))
    lot, nfl = M.lot("P2.mise_en_page", "mise en page")
    out.append(nfl or L.mise_en_page("P2.mise_en_page", gold, lot, M.premier_du_lot))
    out.append(_ligne(M, "P2.input_sheet", "input sheet extraite", ["input_sheet"],
                      lambda isheet: input_sheet("P2.input_sheet", isheet)))
    n = len(ids)
    manquantes = [k for k in range(n + 1) if k not in (iterations_validees or set())]
    out.append(juger("P2.iterations", "iterations validees", n + 1, len(manquantes),
                     f"iterations 0 a {n} ; non validees : {', '.join(map(str, manquantes)) or 'aucune'}"))
    return out + _controles_du_pack(M, 2)


iteration_ = iteration


def porte_3(M, J):
    chemins, nf = M.exiger("P3.certificat", "certificat de la golden", ["certificat", "golden"])
    if nf:
        return [nf]
    cert, gold = chemins
    return [_sur("P3.certificat", "certificat de la golden", lambda: certificat("P3.certificat", cert, gold), [cert, gold])]         + _controles_du_pack(M, 3)


def porte_4(M, J):
    chemins, nf = M.exiger("P4.golden_100", "la golden re-note 100 %", ["golden"])
    if nf:
        return [nf]
    gold = chemins[0]
    out = [_ligne(M, "P4.prompt", "contrat du prompt", ["prompt"], lambda pr: L.verifier_prompt("P4.prompt", pr)),
           _ligne(M, "P4.rubric", "contrat de la rubric", ["rubric"], lambda ru: L.verifier_rubric("P4.rubric", ru)),
           _ligne(M, "P4.golden_100", "la golden re-note 100 %", ["rubric_source"], lambda so: L.noter("P4.golden_100", gold, so))]
    if L.excel_present():
        out.append(_ligne(M, "P4.amplitude", "chaque boucle notee", ["rubric_source"],
                          lambda so: L.amplitude_boucles("P4.amplitude", gold, so)))
    else:
        out.append(non_fait("P4.amplitude", "chaque boucle notee", "Excel absent : les boucles ne peuvent pas etre basculees"))
    out.append(_ligne(M, "P4.tracabilite", "tracabilite prompt <-> rubric", ["prompt", "rubric", "input_sheet"],
                      lambda pr, ru, isheet: L.tracabilite("P4.tracabilite", pr, ru, isheet)))
    return out + _controles_du_pack(M, 4)


def porte_5(M, J):
    chemins, nf = M.exiger("P5.golden", "golden livree", ["golden"])
    if nf:
        return [nf]
    gold = chemins[0]
    out = [L.formules_sures("P5.formules_sures", gold), L.inputs_morts("P5.inputs_morts", gold),
           _ligne(M, "P5.tracabilite", "tracabilite prompt <-> rubric", ["prompt", "rubric", "input_sheet"],
                  lambda pr, ru, isheet: L.tracabilite("P5.tracabilite", pr, ru, isheet))]
    lot, nfl = M.lot("P5.mise_en_page", "mise en page")
    out.append(nfl or L.mise_en_page("P5.mise_en_page", gold, lot, M.premier_du_lot))
    out.append(L.parite_libreoffice("P5.parite", gold))
    out.append(L.gs_audit("P5.audit_golden", gold))
    out.append(L.traces_ia("P5.traces_ia", gold))
    out.append(_ligne(M, "P5.equite", "equite des pertes", ["rubric", "ai_output"],
                      lambda ru, ai: L.dossier_equite("P5.equite", ru, gold, ai)))
    bande = tuple(M.d.get("cible_score") or (0, 45))
    out.append(_ligne(M, "P5.score", "score de l'AI Output", ["score_ai_output", "rubric"],
                      lambda sc, ru: score("P5.score", sc, ru, bande, "score de l'AI Output")))
    out.append(paquet("P5.paquet", M))
    out.append(_sur("P5.metadonnees", "metadonnees nettoyees", lambda: metadonnees("P5.metadonnees", M), []))
    out.append(_sur("P5.calcpr", "calcul iteratif arme", lambda: iteration("P5.calcpr", gold), [gold]))
    return out + _controles_du_pack(M, 5)


def porte_juicing(M, J):
    out = []
    chemins, nf = M.exiger("PJ.golden_100", "la golden re-note 100 %", ["golden"])
    if nf:
        return [nf]
    gold = chemins[0]
    chemins, nf = M.exiger("PJ.documents", "livrables du juicing",
                           ["rubric", "rubric_source", "score_adverse", "note_equite"], section="juicing")
    if nf:
        return [nf]
    rubric, source, sc, note = chemins
    out.append(L.noter("PJ.golden_100", gold, source, "la golden re-note 100 % sur la rubric juicee"))
    out.append(_sur("PJ.score_adverse", "score adverse", lambda: score("PJ.score_adverse", sc, rubric, BANDE_JUICING, "score adverse"), [sc, rubric]))
    out.append(_sur("PJ.note_equite", "note d'equite", lambda: note_equite("PJ.note_equite", note, rubric), [note, rubric]))
    return out + _controles_du_pack(M, "juicing")


PORTES = {"1": porte_1, "3": porte_3, "4": porte_4, "5": porte_5, "juicing": porte_juicing}

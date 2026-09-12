# -*- coding: utf-8 -*-
"""Identites imposees de l'exterieur : ce que le modele n'a pas pense a verifier.

Un modele ne se controle que la ou il a pense a se controler, et le test de
mutation montre pourquoi ca ne suffit pas : la plupart de ses ties sont des
identites algebriques, vraies par construction, qui ferment quoi qu'il arrive.
Un tie ne teste quelque chose que si ses deux membres sont calcules par des
chemins INDEPENDANTS.

Ce module fournit ces chemins. Il ne lit aucun controle du classeur : il
recalcule lui-meme, depuis les valeurs, des egalites qu'un auditeur impose de
l'exterieur.

Les quatre familles, choisies parce qu'aucune ne demande de connaitre le modele :

  C1  continuite ouverture/cloture - le solde de cloture d'une periode est le
      solde d'ouverture de la suivante. Aucune configuration : les libelles
      `opening` et `closing` suffisent a apparier les lignes.
  C2  sommes de bloc - une ligne `Total X` vaut la somme des lignes de son bloc.
      INFORMATIF : sur une feuille en grille, un total agrege des cellules d'une
      autre granularite que les lignes physiquement au-dessus de lui, et
      l'heuristique « remonter jusqu'au trou » se trompe de bloc. Les ecarts
      sont donc affiches mais ne comptent pas comme defauts.
  C3  stabilite de signe - une ligne de cout qui devient positive une annee sur
      dix est une erreur de signe, pas une variation.
  C4  ordre de grandeur - une cellule cent fois ses voisines de ligne est une
      unite qui a glisse ou un doigt qui a derape.

Usage :
    python identites.py "Classeur.xlsx" [--json i.json] [--detail]
"""
import argparse
import collections
import json
import os
import re
import statistics
import sys

import noyau
from openpyxl.utils import get_column_letter


def libelle_de(C, feuille, ligne):
    for col in (3, 2, 4, 1):
        v = C.textes.get((feuille, ligne, col))
        if v and not str(v).startswith("="):
            return str(v).strip()
    return ""


def serie(C, feuille, ligne, cols):
    """Les valeurs numeriques d'une ligne, colonne par colonne."""
    out = {}
    for c in cols:
        v = C.valeurs.get((feuille, ligne, c))
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            out[c] = v
    return out


ENTETE_PERIODE = re.compile(r"^FY-?\d+[AE]?$", re.I)


def colonnes_de_periode(C, feuille, _cache={}):
    """Les colonnes de periode d'une feuille, lues sur SA ligne d'en-tete.

    Version precedente : « toute colonne portant au moins vingt nombres ». Elle
    tenait pour des periodes les colonnes d'un registre de comptes ou d'une
    feuille d'hypotheses, ou une ligne juxtapose une part, une elasticite et un
    plafond. Resultat, 47 faux positifs sur 48 : trois quarts d'une ligne
    « positive » et un quart « negatif » n'a aucun sens quand les quatre
    colonnes ne mesurent pas la meme chose.

    On lit donc l'en-tete - FY-2A, FY0, FY5, FY7E - et une feuille qui n'en
    porte pas n'est simplement pas periodique. C'est le contrat de forme du
    prompt, et il fournit gratuitement la bonne definition.
    """
    if feuille in _cache:
        return _cache[feuille]
    par_ligne = collections.defaultdict(list)
    for (f, l, c), v in C.textes.items():
        if f == feuille and isinstance(v, str) and ENTETE_PERIODE.match(v.strip()):
            par_ligne[l].append(c)
    cols = []
    if par_ligne:
        ligne = max(par_ligne, key=lambda l: len(par_ligne[l]))
        if len(par_ligne[ligne]) >= 4:
            cols = sorted(par_ligne[ligne])
    _cache[feuille] = cols
    return cols


def c1_continuite(C, tolerance=1e-06):
    """Cloture de la periode t = ouverture de la periode t+1."""
    ecarts = []
    testes = 0
    for feuille in C.feuilles:
        cols = colonnes_de_periode(C, feuille)
        if len(cols) < 3:
            continue
        lignes = sorted({l for (f, l, c) in C.valeurs if f == feuille})
        ouvertures, clotures = {}, {}
        for l in lignes:
            lab = libelle_de(C, feuille, l).lower()
            if not lab:
                continue
            cle = re.sub(r"\b(opening|closing|open|close)\b", "", lab).strip(" -,:")
            if re.search(r"\bopening\b|\bopening balance\b", lab):
                ouvertures[cle] = l
            elif re.search(r"\bclosing\b|\bclosing balance\b", lab):
                clotures[cle] = l
        for cle, lo in ouvertures.items():
            lc = clotures.get(cle)
            if lc is None:
                continue
            so, sc = serie(C, feuille, lo, cols), serie(C, feuille, lc, cols)
            for i in range(len(cols) - 1):
                a, b = cols[i], cols[i + 1]
                if a not in sc or b not in so:
                    continue
                testes += 1
                d = abs(sc[a] - so[b])
                base = max(abs(sc[a]), abs(so[b]), 1.0)
                if d / base > tolerance:
                    ecarts.append({
                        "identite": "continuite ouverture/cloture",
                        "feuille": feuille, "poste": cle[:50],
                        "detail": f"cloture {get_column_letter(a)}{lc}={sc[a]:,.2f} "
                                  f"contre ouverture {get_column_letter(b)}{lo}={so[b]:,.2f}",
                        "ecart": sc[a] - so[b]})
    return "C1", "continuite ouverture / cloture", testes, ecarts


def c2_sommes_de_bloc(C, tolerance=1e-06):
    """Une ligne `Total X` vaut la somme des lignes numeriques qui la precedent,
    jusqu'au dernier total ou au dernier trou."""
    ecarts = []
    testes = 0
    for feuille in C.feuilles:
        cols = colonnes_de_periode(C, feuille)
        if len(cols) < 3:
            continue
        lignes = sorted({l for (f, l, c) in C.valeurs if f == feuille})
        ensemble = set(lignes)
        for l in lignes:
            lab = libelle_de(C, feuille, l).lower()
            if not re.match(r"^(total|sum of|sub-?total)\b", lab):
                continue
            # remonter jusqu'au trou ou au total precedent
            membres = []
            ll = l - 1
            while ll in ensemble and len(membres) < 40:
                pl = libelle_de(C, feuille, ll).lower()
                if re.match(r"^(total|sum of|sub-?total)\b", pl):
                    break
                if serie(C, feuille, ll, cols):
                    membres.append(ll)
                ll -= 1
            if len(membres) < 2:
                continue
            st = serie(C, feuille, l, cols)
            for c in cols:
                if c not in st:
                    continue
                somme = sum(C.valeurs.get((feuille, m, c), 0) or 0
                            for m in membres
                            if isinstance(C.valeurs.get((feuille, m, c)), (int, float)))
                testes += 1
                base = max(abs(st[c]), abs(somme), 1.0)
                if abs(st[c] - somme) / base > tolerance:
                    ecarts.append({
                        "identite": "somme de bloc",
                        "feuille": feuille, "poste": lab[:50],
                        "detail": f"{get_column_letter(c)}{l} = {st[c]:,.2f} "
                                  f"contre {len(membres)} lignes qui somment a {somme:,.2f}",
                        "ecart": st[c] - somme})
                break     # une colonne suffit a lever le lievre
    return "C2", "sommes de bloc", testes, ecarts


def c3_signe(C):
    """Une ligne dont le signe bascule sur une periode et une seule."""
    ecarts = []
    testes = 0
    for feuille in C.feuilles:
        cols = colonnes_de_periode(C, feuille)
        if len(cols) < 4:
            continue
        for l in sorted({l for (f, l, c) in C.valeurs if f == feuille}):
            s = serie(C, feuille, l, cols)
            vals = [v for v in s.values() if abs(v) > 1e-09]
            if len(vals) < 4:
                continue
            testes += 1
            pos = sum(1 for v in vals if v > 0)
            neg = len(vals) - pos
            if min(pos, neg) == 1 and max(pos, neg) >= 3:
                deviant = next(c for c, v in s.items()
                               if abs(v) > 1e-09 and (v > 0) == (pos == 1))
                ecarts.append({
                    "identite": "stabilite de signe",
                    "feuille": feuille,
                    "poste": libelle_de(C, feuille, l)[:50],
                    "detail": f"{get_column_letter(deviant)}{l} = {s[deviant]:,.4g} "
                              f"seule de son signe sur {len(vals)} periodes",
                    "ecart": s[deviant]})
    return "C3", "stabilite de signe sur une ligne", testes, ecarts


def c4_ordre_de_grandeur(C, facteur=100.0):
    """Une cellule cent fois la mediane de sa propre ligne."""
    ecarts = []
    testes = 0
    for feuille in C.feuilles:
        cols = colonnes_de_periode(C, feuille)
        if len(cols) < 4:
            continue
        for l in sorted({l for (f, l, c) in C.valeurs if f == feuille}):
            s = serie(C, feuille, l, cols)
            vals = [abs(v) for v in s.values() if abs(v) > 1e-09]
            if len(vals) < 4:
                continue
            testes += 1
            med = statistics.median(vals)
            if med <= 1e-09:
                continue
            for c, v in s.items():
                if abs(v) > facteur * med:
                    ecarts.append({
                        "identite": "ordre de grandeur",
                        "feuille": feuille,
                        "poste": libelle_de(C, feuille, l)[:50],
                        "detail": f"{get_column_letter(c)}{l} = {v:,.4g} contre une "
                                  f"mediane de ligne de {med:,.4g}",
                        "ecart": v})
                    break
    return "C4", "ordre de grandeur dans une ligne", testes, ecarts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("classeur")
    ap.add_argument("--tolerance", type=float, default=1e-06)
    ap.add_argument("--detail", action="store_true")
    ap.add_argument("--json")
    a = ap.parse_args()

    C = noyau.Classeur(a.classeur)
    print(f"IDENTITES IMPOSEES - {os.path.basename(a.classeur)}")
    print("Aucun controle du classeur n'est lu : tout est recalcule ici.\n")

    familles = [c1_continuite(C, a.tolerance), c2_sommes_de_bloc(C, a.tolerance),
                c3_signe(C), c4_ordre_de_grandeur(C)]
    INFORMATIVES = {"C2"}
    total = 0
    sortie = []
    for cle, lib, testes, ecarts in familles:
        info = cle in INFORMATIVES
        etat = "OK  " if not ecarts else ("info" if info else "FAUT")
        print(f"  [{etat}] {cle}  {lib:<38} {testes:>7,} testes"
              f"   {len(ecarts):>5} ecart(s)"
              f"{'   (informatif, a juger)' if info else ''}")
        if not info:
            total += len(ecarts)
        sortie.append({"cle": cle, "libelle": lib, "testes": testes,
                       "ecarts": ecarts[:200], "nombre": len(ecarts)})
        if a.detail and ecarts:
            par = collections.Counter(e["feuille"] for e in ecarts)
            for f, n in par.most_common(6):
                print(f"          {f:<34} {n:>5}")
            for e in ecarts[:6]:
                print(f"          {e['poste'][:42]:<42} {e['detail'][:74]}")

    print()
    if total == 0:
        print("OK : les quatre familles d'identites tiennent.")
        print("     Ce sont des chemins independants des ties du modele : le")
        print("     classeur ne s'auto-confirme pas, il resiste a un tiers.")
    else:
        print(f"{total} ecart(s) sur des identites que le modele ne declare pas.")
        print("Chacun est un endroit ou le classeur se contredit sans le savoir.")

    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump({"classeur": os.path.basename(a.classeur),
                       "ecarts_total": total, "familles": sortie},
                      fh, ensure_ascii=False, indent=2)
        print(f"\necrit : {a.json}")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())

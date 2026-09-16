"""Les ateliers avec l'auteur (references/ateliers.md) : preparer une copie de travail de la golden, l'ouvrir dans Excel
sur la cellule qui compte, puis relire ce que l'auteur y a change, cellule par cellule.

    python atelier.py preparer "<golden recalculee>" --atelier "<build>/ateliers/GS atelier 03 - X.xlsx"
    python atelier.py ouvrir   "<copie>" --cellule "Input_Sheet!E812"
    python atelier.py relire   "<copie>" --reference "<golden recalculee>" [--suivre "Feuille!A1" ...]
                               [--registre noms.json] [--json sortie.json]
    python atelier.py decider  "<build>/ateliers/it03.json" --cellule "Feuille!A1" --decision reportee|abandonnee

`relire` ne sait rien du pack : le libelle d'une modification est le premier texte de sa ligne a gauche de la cellule,
son en-tete le premier texte de sa colonne au-dessus. `--registre` accepte un JSON {"Feuille!A1": "nom"} pour nommer
les cellules. L'outil rend toujours ce qu'il a compare : un « 0 modification » sans compte ne prouve rien.

`decider` inscrit la decision du owner sur une modification relue, dans le JSON ecrit par `relire --json` : c'est ce
fichier que la porte d'iteration lit (references/portes.md). `--cellule tout` decide toutes celles qui n'en ont pas.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ICI = Path(__file__).resolve().parent


def verrou(chemin: Path) -> Path:
    return chemin.with_name("~$" + chemin.name)


def preparer(a):
    src, dst = Path(a.golden), Path(a.atelier)
    if dst.exists():
        sys.exit(f"refus : {dst} existe deja ; une copie d'atelier ne s'ecrase jamais, numeroter la suivante")
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dst)
    print(f"copie d'atelier : {dst} ({dst.stat().st_size} octets, depuis {src.name})")


def ouvrir(a):
    chemin = Path(a.copie).resolve()
    if verrou(chemin).exists():
        sys.exit(f"deja ouvert dans Excel (verrou {verrou(chemin).name}) : rien a faire")
    args = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ICI / "atelier.ps1"), "-Chemin", str(chemin)]
    if a.cellule:
        args += ["-Cellule", a.cellule]
    r = subprocess.run(args, capture_output=True, text=True)
    print((r.stdout or "").strip(), (r.stderr or "").strip()[-800:])
    sys.exit(r.returncode)


def libelle(ws, r, c):
    for cc in range(c - 1, 0, -1):
        v = ws.cell(r, cc).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()[:70]
    return ""


def entete(ws, r, c):
    for rr in range(r - 1, 0, -1):
        v = ws.cell(rr, c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()[:40]
    return ""


def court(v):
    if isinstance(v, float):
        return round(v, 6)
    if isinstance(v, datetime):
        return v.date().isoformat()
    return v


def relire(a):
    copie, ref = Path(a.copie), Path(a.reference)
    if verrou(copie).exists():
        sys.exit(f"refus : {copie.name} est encore ouvert dans Excel (verrou {verrou(copie).name}) ; le fermer d'abord")
    registre = json.loads(Path(a.registre).read_text(encoding="utf-8")) if a.registre else {}
    wc, wr = load_workbook(copie), load_workbook(ref)
    modifs, comparees = [], 0
    feuilles_c, feuilles_r = set(wc.sheetnames), set(wr.sheetnames)
    for f in sorted(feuilles_c ^ feuilles_r):
        modifs.append(dict(feuille=f, cellule="", nature="onglet ajoute" if f in feuilles_c else "onglet supprime"))
    for f in wr.sheetnames:
        if f not in feuilles_c:
            continue
        sc, sr = wc[f], wr[f]
        nr, nc = max(sc.max_row, sr.max_row), max(sc.max_column, sr.max_column)
        for r in range(1, nr + 1):
            for c in range(1, nc + 1):
                vc, vr = sc.cell(r, c).value, sr.cell(r, c).value
                if vc is None and vr is None:
                    continue
                comparees += 1
                if vc == vr or (isinstance(vc, (int, float)) and isinstance(vr, (int, float)) and abs(vc - vr) < 1e-12):
                    continue
                formule = any(isinstance(x, str) and x.startswith("=") for x in (vc, vr))
                nature = "ajout" if vr is None else "effacement" if vc is None else "formule" if formule else "entree"
                adr = f"{f}!{get_column_letter(c)}{r}"
                base = sc if vc is not None else sr
                modifs.append(dict(feuille=f, cellule=f"{get_column_letter(c)}{r}", nom=registre.get(adr, ""),
                                   libelle=libelle(base, r, c), entete=entete(base, r, c), nature=nature,
                                   avant=court(vr), apres=court(vc)))
    suivis = []
    if a.suivre:
        vc, vr = load_workbook(copie, data_only=True), load_workbook(ref, data_only=True)
        for s in a.suivre:
            f, cel = s.rsplit("!", 1)
            f = f.strip("'")
            avant = vr[f][cel].value if f in vr.sheetnames else None
            apres = vc[f][cel].value if f in vc.sheetnames else None
            suivis.append(dict(cellule=s, libelle=libelle(wr[f], wr[f][cel].row, wr[f][cel].column) if f in wr.sheetnames else "",
                               avant=court(avant), apres=court(apres)))
    print(f"{comparees} cellules non vides comparees sur {len(wr.sheetnames)} onglets ; {len(modifs)} modification(s)")
    for m in modifs:
        if not m["cellule"]:
            print(f"  [{m['nature']}] {m['feuille']}")
            continue
        nom = f" ({m['nom']})" if m["nom"] else ""
        print(f"  [{m['nature']}] {m['feuille']}!{m['cellule']}{nom} | {m['libelle']} | {m['entete']} | {m['avant']!r} -> {m['apres']!r}")
    for s in suivis:
        print(f"  [suivi] {s['cellule']} | {s['libelle']} | {s['avant']!r} -> {s['apres']!r}")
    if a.json:
        Path(a.json).write_text(json.dumps(dict(comparees=comparees, modifications=modifs, suivis=suivis), ensure_ascii=False,
                                           indent=1, default=str), encoding="utf-8")
    if comparees == 0:
        sys.exit("aucune cellule comparee : la relecture n'a rien lu")


def decider(a):
    chemin = Path(a.fichier)
    d = json.loads(chemin.read_text(encoding="utf-8"))
    cible = a.cellule.replace("$", "")
    touchees = 0
    for m in d.get("modifications", []):
        if not m.get("cellule"):
            continue
        adr = f"{m['feuille']}!{m['cellule']}"
        if (cible == "tout" and not m.get("decision")) or adr == cible or adr == cible.replace("'", ""):
            m["decision"] = a.decision
            touchees += 1
    if touchees == 0:
        sys.exit(f"refus : aucune modification {a.cellule!r} dans {chemin.name}")
    chemin.write_text(json.dumps(d, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    restantes = sum(1 for m in d.get("modifications", []) if m.get("cellule") and not m.get("decision"))
    print(f"{touchees} modification(s) decidee(s) {a.decision} ; {restantes} sans decision")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("preparer")
    p.add_argument("golden")
    p.add_argument("--atelier", required=True)
    p.set_defaults(fn=preparer)
    p = sub.add_parser("ouvrir")
    p.add_argument("copie")
    p.add_argument("--cellule", default="")
    p.set_defaults(fn=ouvrir)
    p = sub.add_parser("relire")
    p.add_argument("copie")
    p.add_argument("--reference", required=True)
    p.add_argument("--suivre", action="append")
    p.add_argument("--registre")
    p.add_argument("--json")
    p.set_defaults(fn=relire)
    p = sub.add_parser("decider")
    p.add_argument("fichier")
    p.add_argument("--cellule", required=True)
    p.add_argument("--decision", required=True, choices=("reportee", "abandonnee"))
    p.set_defaults(fn=decider)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
